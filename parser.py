import sys

# ──────────────────────────────────────────────
# Проверка версии Python при старте
# ──────────────────────────────────────────────
MIN_PYTHON = (3, 10)
if sys.version_info < MIN_PYTHON:
    sys.exit(
        f"❌ Требуется Python {MIN_PYTHON[0]}.{MIN_PYTHON[1]}+, "
        f"а запущен {sys.version_info.major}.{sys.version_info.minor}.\n"
        f"Установите актуальную версию и пересоздайте .venv."
    )

# Проверка наличия зависимостей
try:
    import openpyxl
    import requests
    from bs4 import BeautifulSoup
except ImportError as e:
    sys.exit(
        f"❌ Не установлена зависимость: {e}\n"
        f"Активируйте .venv и выполните: pip install -r requirements.txt"
    )

import re
import time

# ==========================================
# НАСТРОЙКИ
# ==========================================
INPUT_FILE = 'MUN_Zadanie_2026_AVGUST.xlsx'
OUTPUT_FILE = 'MUN_Zadanie_2026_AVGUST_ready.xlsx'
TARGET_SHEET_KEYWORD = 'сентябрь'

HEADERS = {
    'User-Agent': ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) '
                   'AppleWebKit/537.36 (KHTML, like Gecko) '
                   'Chrome/125.0.0.0 Safari/537.36'),
    'Accept-Language': 'ru-RU,ru;q=0.9',
}

session = requests.Session()
session.headers.update(HEADERS)


def get_url_from_row(row_cells):
    """Ищет URL в строке: сначала гиперссылки, потом текст."""
    for cell in row_cells:
        if cell.hyperlink and cell.hyperlink.target:
            target = str(cell.hyperlink.target).strip()
            if target.startswith('http') or 'sgpress.ru' in target:
                return target
    for cell in row_cells:
        val = str(cell.value).strip() if cell.value else ""
        m = re.search(r'https?://\S+', val)
        if match := m:
            return match.group(0).rstrip('|').strip()
        if 'sgpress.ru/news/' in val:
            return 'https://' + val if not val.startswith('http') else val
    return None


def count_chars_no_spaces(url):
    """Считает знаки без пробелов от заголовка до маркера конца статьи."""
    try:
        if not url.startswith('http'):
            url = 'https://' + url

        resp = session.get(url, timeout=20)
        resp.encoding = 'utf-8'

        if resp.status_code != 200:
            return f"HTTP {resp.status_code}"

        soup = BeautifulSoup(resp.text, 'html.parser')

        for tag in soup(['script', 'style', 'noscript']):
            tag.decompose()

        # Ищем основной контент
        article_text = ""
        for sel in ['.contant-block-article', 'article', '.entry-content',
                    '.post-content', 'main']:
            block = soup.select_one(sel)
            if block and len(block.get_text(strip=True)) > 200:
                article_text = block.get_text(separator=' ', strip=True)
                break

        if not article_text:
            article_text = soup.get_text(separator=' ', strip=True)

        # Обрезаем по маркеру конца
        for marker in ["Читай, где удобно", "Читайте, где удобно",
                       "Читай нас", "Читайте нас", "Подписывайтесь на"]:
            idx = article_text.find(marker)
            if idx != -1:
                article_text = article_text[:idx]
                break

        return len(re.sub(r'\s+', '', article_text))

    except Exception as e:
        return f"Ошибка: {str(e)[:40]}"


def main():
    if not os.path.exists(INPUT_FILE):
        print(f"❌ Файл '{INPUT_FILE}' не найден в папке со скриптом.")
        return

    print(f"⏳ Загрузка '{INPUT_FILE}'...")
    wb = openpyxl.load_workbook(INPUT_FILE)

    sheet_name = None
    for name in wb.sheetnames:
        if TARGET_SHEET_KEYWORD in name.lower():
            sheet_name = name
            break

    if not sheet_name:
        print(f"❌ Вкладка '{TARGET_SHEET_KEYWORD}' не найдена. "
              f"Доступные: {wb.sheetnames}")
        return

    sheet = wb[sheet_name]
    print(f"✅ Вкладка: '{sheet_name}'")

    # Ищем заголовки
    header_row = url_col = chars_col = None
    for row_idx in range(1, 30):
        for cell in sheet[row_idx]:
            val = str(cell.value).strip() if cell.value else ""
            if 'ссылка' in val.lower():
                url_col = cell.column
                header_row = row_idx
            if 'площадь' in val.lower() or 'знаков' in val.lower():
                chars_col = cell.column

    if not all([header_row, url_col, chars_col]):
        print("❌ Не найдены заголовки 'Ссылка' и 'Площадь опубликования'.")
        return

    print(f"✅ Заголовки в строке {header_row}: "
          f"Ссылка=столбец {url_col}, Знаки=столбец {chars_col}")
    print("\n🔍 Начинаем обработку...\n")

    processed = skipped = empty = errors = 0

    for r in range(header_row + 1, sheet.max_row + 1):
        row_cells = sheet[r]
        url = get_url_from_row(row_cells)

        if not url:
            empty += 1
            continue

        chars_cell = sheet.cell(row=r, column=chars_col)
        cur = str(chars_cell.value).strip() if chars_cell.value else ""
        try:
            if float(cur) > 0:
                skipped += 1
                continue
        except ValueError:
            pass

        print(f"Строка {r}: {url[:70]}...")
        result = count_chars_no_spaces(url)
        chars_cell.value = result

        if isinstance(result, int):
            print(f"  ✅ {result} знаков")
            processed += 1
        else:
            print(f"  ❌ {result}")
            errors += 1

        time.sleep(1.5)

    wb.save(OUTPUT_FILE)
    print(f"\n{'='*50}")
    print(f"🏁 Готово!")
    print(f"   ✅ Обработано: {processed}")
    print(f"   ⏭️ Пропущено: {skipped}")
    print(f"   ⚠️ Без ссылок: {empty}")
    print(f"   ❌ Ошибок: {errors}")
    print(f"💾 Сохранено в '{OUTPUT_FILE}'")


import os

if __name__ == "__main__":
    main()